"""
scripts/artifacts.py
Uploaded career artifacts under workspace/uploads/: saving, listing, and plain-text extraction
(PDF via pypdf, DOCX via python-docx, HTML via htmltext, everything else read as text).
Extractions are cached in workspace/uploads/.text/.
"""

import re
from pathlib import Path
from typing import Dict, List, Tuple

from config import UPLOADS_DIR
from htmltext import html_to_text

ALLOWED_SUFFIXES = {".pdf", ".docx", ".md", ".txt", ".html", ".htm", ".csv", ".json"}
TEXT_DIR = UPLOADS_DIR / ".text"


def safe_name(name: str) -> str:
    base = Path(name).name
    stem, suffix = Path(base).stem, Path(base).suffix.lower()
    stem = re.sub(r"[^A-Za-z0-9._ -]+", "_", stem).strip(" .") or "upload"
    return f"{stem[:120]}{suffix}"


def extract_text(path: Path) -> str:
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        from pypdf import PdfReader
        return "\n".join((page.extract_text() or "") for page in PdfReader(str(path)).pages)
    if suffix == ".docx":
        import docx
        doc = docx.Document(str(path))
        parts = [p.text for p in doc.paragraphs]
        for table in doc.tables:
            for row in table.rows:
                parts.append(" | ".join(cell.text.strip() for cell in row.cells))
        return "\n".join(parts)
    raw = path.read_text(encoding="utf-8", errors="ignore")
    return html_to_text(raw) if suffix in (".html", ".htm") else raw


def cached_text(path: Path) -> str:
    cache = TEXT_DIR / f"{path.name}.txt"
    if cache.exists() and cache.stat().st_mtime >= path.stat().st_mtime:
        return cache.read_text(encoding="utf-8")
    text = re.sub(r"\n{3,}", "\n\n", extract_text(path)).strip()
    TEXT_DIR.mkdir(parents=True, exist_ok=True)
    cache.write_text(text, encoding="utf-8")
    return text


def list_uploads() -> List[Dict]:
    if not UPLOADS_DIR.exists():
        return []
    out = []
    for p in sorted(UPLOADS_DIR.iterdir()):
        if not p.is_file() or p.name.startswith(".") or p.suffix.lower() not in ALLOWED_SUFFIXES:
            continue
        cache = TEXT_DIR / f"{p.name}.txt"
        chars = len(cache.read_text(encoding="utf-8")) if cache.exists() else None
        out.append({"name": p.name, "size": p.stat().st_size, "kind": p.suffix.lower().lstrip("."), "extracted_chars": chars})
    return out


def save_upload(name: str, data: bytes) -> Path:
    fname = safe_name(name)
    if Path(fname).suffix not in ALLOWED_SUFFIXES:
        raise ValueError(f"Unsupported file type: {Path(name).suffix or 'none'}. Allowed: {', '.join(sorted(ALLOWED_SUFFIXES))}")
    UPLOADS_DIR.mkdir(parents=True, exist_ok=True)
    dest = UPLOADS_DIR / fname
    dest.write_bytes(data)
    return dest


def delete_upload(name: str) -> None:
    fname = Path(name).name
    for p in (UPLOADS_DIR / fname, TEXT_DIR / f"{fname}.txt"):
        if p.is_file():
            p.unlink()


def all_texts(max_total: int = 250_000, max_each: int = 80_000) -> Tuple[List[Tuple[str, str]], List[str]]:
    """Returns ([(filename, text)], notes) within the character budget, extracting as needed."""
    texts: List[Tuple[str, str]] = []
    notes: List[str] = []
    budget = max_total
    for item in list_uploads():
        path = UPLOADS_DIR / item["name"]
        try:
            text = cached_text(path)
        except Exception as e:
            notes.append(f"{item['name']}: could not read ({e})")
            continue
        if not text.strip():
            notes.append(f"{item['name']}: no extractable text (scanned PDF?)")
            continue
        clipped = text[:min(max_each, budget)]
        if len(clipped) < len(text):
            notes.append(f"{item['name']}: truncated to {len(clipped):,} of {len(text):,} chars")
        if clipped:
            texts.append((item["name"], clipped))
            budget -= len(clipped)
        if budget <= 0:
            notes.append("Character budget reached; remaining uploads skipped.")
            break
    return texts, notes
